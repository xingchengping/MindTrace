"""文档文本提取：PDF / Word / Excel（低频深采，运行于 watch_dirs）。"""
from __future__ import annotations

import logging
from pathlib import Path

log = logging.getLogger("mindtrace.collector.docs")

_TEXT_CAP = 2000
_EXTS = (".pdf", ".docx", ".xlsx")


def extract_text(path: Path, max_chars: int = _TEXT_CAP) -> str | None:
    ext = path.suffix.lower()
    try:
        if ext == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(str(path))
            texts = []
            for page in reader.pages[:5]:
                t = page.extract_text() or ""
                texts.append(t)
                if sum(len(x) for x in texts) > max_chars:
                    break
            return "\n".join(texts)[:max_chars]
        if ext == ".docx":
            import docx

            doc = docx.Document(str(path))
            return "\n".join(p.text for p in doc.paragraphs[:60])[:max_chars]
        if ext == ".xlsx":
            import openpyxl

            wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
            out = []
            for ws in wb.worksheets[:2]:
                for row in ws.iter_rows(max_row=30, values_only=True):
                    vals = [str(c) for c in row if c is not None]
                    if vals:
                        out.append(" | ".join(vals))
                if sum(len(x) for x in out) > max_chars:
                    break
            wb.close()
            return "\n".join(out)[:max_chars]
    except Exception as exc:  # pragma: no cover
        log.debug("文档提取失败 %s: %s", path.name, exc)
    return None


def scan_recent_docs(dirs: list[str], max_docs: int = 5) -> list[dict]:
    """扫描工作目录中最近 1 小时修改的文档，提取文本。"""
    from datetime import datetime, timedelta

    results = []
    for d in dirs:
        root = Path(d)
        if not root.exists():
            continue
        since = datetime.now() - timedelta(hours=1)
        try:
            candidates = sorted(
                (p for p in root.rglob("*") if p.suffix.lower() in _EXTS),
                key=lambda p: p.stat().st_mtime, reverse=True,
            )
        except Exception:
            candidates = []
        for p in candidates[:max_docs]:
            try:
                mtime = datetime.fromtimestamp(p.stat().st_mtime)
            except Exception:
                continue
            if mtime < since:
                continue
            text = extract_text(p)
            if text:
                results.append({"path": str(p), "name": p.name, "text": text, "time": mtime})
    return results
